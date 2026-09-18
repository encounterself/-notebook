import json, re, math, sys
from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook

path=Path(sys.argv[1])
def norm(x):
    if x is None:return ""
    return re.sub(r"[^a-z0-9]+","_",str(x).replace("\u00a0"," ").strip().lower()).strip("_")
def num(v):
    if v is None or isinstance(v,bool):return None
    if isinstance(v,(int,float)):
        return float(v) if math.isfinite(float(v)) else None
    s=str(v).strip().replace(",","").replace("$","")
    if s in ("","-","—","n/a","na","null"):return None
    neg=s.startswith("(") and s.endswith(")")
    if neg:s=s[1:-1]
    try:
        x=float(s);return -x if neg else x
    except:return None
def dt(v):
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    if v is None:return None
    for f in ("%Y-%m-%d","%Y/%m/%d","%m/%d/%Y","%m/%d/%y","%d/%m/%Y"):
        try:return datetime.strptime(str(v).strip(),f).date()
        except:pass
    return None
def header(ws):
    for i,row in enumerate(ws.iter_rows(values_only=True),1):
        vals=list(row)
        if i<=20 and sum(v not in (None,"") for v in vals)>=2:
            return i,[str(v).strip() if v is not None else "" for v in vals]
        if i>=20:break
    return 1,[]
amounts={"final_charge","final_charge_for_usage_days","pending_charge","credit_owed","monthly_rate","price","final_price","correct_price","original_charge","total_charge","amount"}
dates={"final_start_date","final_end_date","cycle_start_date","cycle_end_date","cycle_start","cycle_end","new_activation_date","charge_start_date","charge_end_date","charge_start","charge_end","offstock_date","date_line_went_down","date_replacement_was_activated"}
keys={"imsi","iccid","charge_type","product_name","prod","source_file","sheet_name","final_days","active_days","usage_days"}
item={"local_excel_file":str(path),"file_name":path.name,"extension":path.suffix.lower(),"file_size":path.stat().st_size,"is_temp_lock_file":path.name.startswith("~$"),"read_status":"OK","sheets":[]}
try:
    wb=load_workbook(path,read_only=True,data_only=True)
    item["sheet_names"]=wb.sheetnames
    for sn in wb.sheetnames:
        ws=wb[sn];hr,heads=header(ws);nh=[norm(h) for h in heads]
        ai={j:heads[j] for j,h in enumerate(nh) if h in amounts or any(h.endswith("_"+a) for a in amounts)}
        di={j:heads[j] for j,h in enumerate(nh) if h in dates or any(h.endswith("_"+a) for a in dates)}
        ast={h:{"count_numeric":0,"sum":0.0,"min":None,"max":None} for h in ai.values()}
        dst={h:{"count_date":0,"min":None,"max":None} for h in di.values()}
        rows=0;nonempty=0
        for rn,row in enumerate(ws.iter_rows(values_only=True),1):
            if rn<=hr:continue
            vals=list(row);rows+=1
            if any(v not in (None,"") for v in vals):nonempty+=1
            for j,h in ai.items():
                v=num(vals[j] if j<len(vals) else None)
                if v is not None:
                    z=ast[h];z["count_numeric"]+=1;z["sum"]+=v;z["min"]=v if z["min"] is None else min(z["min"],v);z["max"]=v if z["max"] is None else max(z["max"],v)
            for j,h in di.items():
                v=dt(vals[j] if j<len(vals) else None)
                if v is not None:
                    z=dst[h];sv=str(v);z["count_date"]+=1;z["min"]=sv if z["min"] is None else min(z["min"],sv);z["max"]=sv if z["max"] is None else max(z["max"],sv)
        item["sheets"].append({"sheet_name":sn,"header_row":hr,"max_row":ws.max_row,"max_column":ws.max_column,"headers":heads,"key_columns":[heads[j] for j,h in enumerate(nh) if h in keys or any(h.endswith("_"+k) for k in keys)],"data_rows_after_header":rows,"nonempty_data_rows":nonempty,"amount_stats":ast,"date_stats":dst})
    wb.close()
except Exception as e:
    item["read_status"]="INVALID_OR_UNREADABLE";item["sheet_names"]=[];item["sheets"]=[];item["error"]=f"{type(e).__name__}: {e}"
print(json.dumps(item,ensure_ascii=False))


