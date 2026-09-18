import csv
import os
import re
from openpyxl import load_workbook

ROOT = r"C:\Users\EDY\Desktop\WA账单20260914"
OUT = r"C:\Users\EDY\Desktop\对账notebook\recon\agent_runs\A\excel_control_A_local_formula_scan.tsv"
rows = []
for dp, _, fs in os.walk(ROOT):
    for name in sorted(fs):
        if not name.lower().endswith(".xlsx") or name.startswith("~$"):
            continue
        path = os.path.join(dp, name)
        try:
            wb = load_workbook(path, read_only=True, data_only=False)
            for ws in wb.worksheets:
                formula_count = 0
                errors = []
                samples = []
                for row in ws.iter_rows():
                    for cell in row:
                        value = cell.value
                        if isinstance(value, str) and value.startswith("="):
                            formula_count += 1
                            if len(samples) < 5:
                                samples.append(f"{cell.coordinate}:{value}")
                            for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"):
                                if token in value.upper():
                                    errors.append(token)
                rows.append({
                    "local_excel_file": name,
                    "local_relative_path": os.path.relpath(path, ROOT),
                    "sheet_name": ws.title,
                    "formula_count": formula_count,
                    "formula_error_count": len(errors),
                    "formula_error_tokens": "|".join(sorted(set(errors))),
                    "formula_samples": "|".join(samples),
                    "formula_status": "FORMULA_TEXT_HAS_ERROR_TOKEN" if errors else "NO_ERROR_TOKEN_IN_FORMULA_TEXT",
                })
            wb.close()
        except Exception as exc:
            rows.append({
                "local_excel_file": name,
                "local_relative_path": os.path.relpath(path, ROOT),
                "sheet_name": "",
                "formula_count": "",
                "formula_error_count": "",
                "formula_error_tokens": "",
                "formula_samples": "",
                "formula_status": "READ_ERROR:" + type(exc).__name__,
            })
fields = ["local_excel_file","local_relative_path","sheet_name","formula_count","formula_error_count","formula_error_tokens","formula_samples","formula_status"]
with open(OUT, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=fields, delimiter="\t")
    w.writeheader()
    w.writerows(rows)

