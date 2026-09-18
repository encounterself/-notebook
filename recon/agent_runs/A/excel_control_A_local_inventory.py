import csv
import json
import math
import os
import re
import unicodedata
import traceback
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

SOURCE_ROOT = r"C:\Users\EDY\Desktop\WA账单20260914"
OUT_DIR = r"C:\Users\EDY\Desktop\对账notebook\recon\agent_runs\A"

def nfkc(value):
    return unicodedata.normalize("NFKC", str(value)).replace("\u00a0", " ")

def canon(value):
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "", nfkc(value).strip().lower())

def raw_text(value):
    if value is None:
        return ""
    return str(value).replace(chr(9), " ").replace(chr(13), " ").replace(chr(10), " ").strip()
def safe_text(value):
    if value is None:
        return ""
    return nfkc(value).replace("\t", " ").replace("\r", " ").replace("\n", " ").strip()

def parse_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return float(value)
    text = safe_text(value).replace(",", "").replace("$", "").replace("USD", "").strip()
    if not text:
        return None
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        return None

def parse_date(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = safe_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text[:19], fmt).date()
        except ValueError:
            pass
    return None

def json_compact(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))

def header_score(row):
    known = {
        "imsi", "iccid", "productname", "chargetype", "sourcefile", "sheetname",
        "finalstartdate", "finalenddate", "finaldays", "activedays", "usagedays",
        "usagedaysalt", "usagedaysgtactive", "monthlyrate", "price", "finalcharge",
        "finalchargeforusagedays", "creditowed", "finalprice", "correctprice",
        "originalcharge", "pendingcharge", "cyclestart", "cycleend", "cyclestartdate",
        "cycleenddate", "datelinewentdown", "datereplacementwasactivated",
        "dayswithoutservice", "status", "newactivationdate", "offstockdate",
    }
    return sum(1 for item in row if canon(item) in known)

def choose_header(ws):
    best_idx = 1
    best_score = -1
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 20), values_only=True), start=1):
        score = header_score(list(row))
        if score > best_score:
            best_idx = idx
            best_score = score
    if best_score <= 0:
        best_idx = 1
    return best_idx, best_score

def classify_field(header):
    c = canon(header)
    if c == "imsi":
        return "imsi"
    if c == "chargetype":
        return "charge_type"
    if c == "sourcefile":
        return "source_file"
    amount_names = {
        "finalcharge", "finalchargeforusagedays", "originalcharge", "pendingcharge",
        "creditowed", "price", "finalprice", "monthlyrate", "correctprice", "amount",
    }
    if c in amount_names or "charge" in c or "price" in c or "amount" in c:
        return "amount"
    date_names = {
        "finalstartdate", "finalenddate", "newactivationdate", "offstockdate",
        "cyclestart", "cycleend", "cyclestartdate", "cycleenddate",
        "datelinewentdown", "datereplacementwasactivated",
    }
    if c in date_names or "date" in c or c.startswith("cycle"):
        return "date"
    if "day" in c:
        return "days"
    if c in {"iccid", "productname", "prod", "replacementproduct", "status"}:
        return "key_or_label"
    return "other"

def write_tsv(path, fieldnames, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: (raw_text(row.get(key, "")) if key in {"local_excel_file", "local_relative_path", "local_full_path"} else safe_text(row.get(key, ""))) for key in fieldnames})

def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    all_files = []
    for root, _, names in os.walk(SOURCE_ROOT):
        for name in names:
            path = os.path.join(root, name)
            ext = os.path.splitext(name)[1].lower()
            all_files.append({
                "local_excel_file": name,
                "local_relative_path": os.path.relpath(path, SOURCE_ROOT),
                "local_full_path": path,
                "extension": ext,
                "file_size_bytes": os.path.getsize(path),
            })
    all_files.sort(key=lambda row: row["local_relative_path"].casefold())

    file_rows = []
    sheet_rows = []
    field_rows = []

    for file_info in all_files:
        path = file_info["local_full_path"]
        ext = file_info["extension"]
        base = file_info["local_excel_file"]
        if ext not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            file_rows.append({**file_info, "file_class": "non_excel_file", "workbook_status": "NOT_INSPECTED"})
            continue
        if base.startswith("~$"):
            file_rows.append({**file_info, "file_class": "excel_lock_file", "workbook_status": "EXCLUDED_TEMPORARY_LOCK_FILE"})
            continue
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            file_rows.append({**file_info, "file_class": "excel_source_candidate", "workbook_status": "READ_ERROR: " + type(exc).__name__})
            continue

        sheet_names = list(wb.sheetnames)
        file_rows.append({
            **file_info,
            "file_class": "excel_source_candidate",
            "workbook_status": "READ_OK",
            "sheet_names": "|".join(sheet_names),
            "sheet_count": len(sheet_names),
        })

        for ws in wb.worksheets:
            header_row_num, header_score_value = choose_header(ws)
            header_values = next(ws.iter_rows(min_row=header_row_num, max_row=header_row_num, values_only=True), ())
            headers = [safe_text(value) or f"__blank_{idx+1}" for idx, value in enumerate(header_values)]
            while headers and headers[-1].startswith("__blank_"):
                headers.pop()
            field_kind = {idx: classify_field(header) for idx, header in enumerate(headers)}
            stats = {
                "row_count": 0,
                "imsi_values": set(),
                "amount": defaultdict(lambda: {"count": 0, "sum": 0.0, "min": None, "max": None}),
                "date": defaultdict(lambda: {"count": 0, "min": None, "max": None}),
                "source_file_values": set(),
                "charge_type_values": set(),
            }

            for values in ws.iter_rows(min_row=header_row_num + 1, values_only=True):
                row_values = list(values[:len(headers)])
                if not any(value is not None and safe_text(value) != "" for value in row_values):
                    continue
                stats["row_count"] += 1
                for idx, header in enumerate(headers):
                    value = row_values[idx] if idx < len(row_values) else None
                    kind = field_kind[idx]
                    if kind == "imsi" and safe_text(value):
                        stats["imsi_values"].add(safe_text(value))
                    elif kind == "source_file" and safe_text(value):
                        stats["source_file_values"].add(safe_text(value))
                    elif kind == "charge_type" and safe_text(value):
                        stats["charge_type_values"].add(safe_text(value))
                    elif kind == "amount":
                        number = parse_number(value)
                        if number is not None:
                            current = stats["amount"][header]
                            current["count"] += 1
                            current["sum"] += number
                            current["min"] = number if current["min"] is None else min(current["min"], number)
                            current["max"] = number if current["max"] is None else max(current["max"], number)
                    elif kind == "date":
                        parsed = parse_date(value)
                        if parsed is not None:
                            current = stats["date"][header]
                            current["count"] += 1
                            current["min"] = parsed.isoformat() if current["min"] is None else min(current["min"], parsed.isoformat())
                            current["max"] = parsed.isoformat() if current["max"] is None else max(current["max"], parsed.isoformat())

            amount_summary = {
                header: {
                    "count": metric["count"],
                    "sum": round(metric["sum"], 9),
                    "min": metric["min"],
                    "max": metric["max"],
                }
                for header, metric in sorted(stats["amount"].items())
            }
            date_summary = {header: metric for header, metric in sorted(stats["date"].items())}
            key_columns = [header for idx, header in enumerate(headers) if field_kind[idx] in {"imsi", "charge_type", "source_file", "amount", "date", "days", "key_or_label"}]
            sheet_rows.append({
                "local_excel_file": base,
                "local_relative_path": file_info["local_relative_path"],
                "sheet_name": ws.title,
                "header_row": header_row_num,
                "header_detection_score": header_score_value,
                "data_row_count": stats["row_count"],
                "worksheet_max_row": ws.max_row,
                "worksheet_max_column": ws.max_column,
                "column_names": "|".join(headers),
                "key_columns": "|".join(key_columns),
                "imsi_count": len(stats["imsi_values"]),
                "charge_type_values": "|".join(sorted(stats["charge_type_values"])),
                "source_file_values": "|".join(sorted(stats["source_file_values"])),
                "amount_columns": "|".join(sorted(stats["amount"])),
                "amount_summary_json": json_compact(amount_summary),
                "date_columns": "|".join(sorted(stats["date"])),
                "date_summary_json": json_compact(date_summary),
            })
            for header, metric in sorted(stats["amount"].items()):
                field_rows.append({
                    "local_excel_file": base,
                    "local_relative_path": file_info["local_relative_path"],
                    "sheet_name": ws.title,
                    "field_name": header,
                    "field_kind": "amount",
                    "non_null_count": metric["count"],
                    "sum_value": round(metric["sum"], 9),
                    "min_value": metric["min"],
                    "max_value": metric["max"],
                })
            for header, metric in sorted(stats["date"].items()):
                field_rows.append({
                    "local_excel_file": base,
                    "local_relative_path": file_info["local_relative_path"],
                    "sheet_name": ws.title,
                    "field_name": header,
                    "field_kind": "date",
                    "non_null_count": metric["count"],
                    "min_value": metric["min"],
                    "max_value": metric["max"],
                })
        wb.close()

    write_tsv(os.path.join(OUT_DIR, "excel_control_A_local_file_inventory.tsv"), [
        "local_excel_file", "local_relative_path", "local_full_path", "extension", "file_size_bytes",
        "file_class", "workbook_status", "sheet_names", "sheet_count",
    ], file_rows)
    write_tsv(os.path.join(OUT_DIR, "excel_control_A_local_sheet_inventory.tsv"), [
        "local_excel_file", "local_relative_path", "sheet_name", "header_row", "header_detection_score",
        "data_row_count", "worksheet_max_row", "worksheet_max_column", "column_names", "key_columns",
        "imsi_count", "charge_type_values", "source_file_values", "amount_columns", "amount_summary_json",
        "date_columns", "date_summary_json",
    ], sheet_rows)
    write_tsv(os.path.join(OUT_DIR, "excel_control_A_local_field_stats.tsv"), [
        "local_excel_file", "local_relative_path", "sheet_name", "field_name", "field_kind",
        "non_null_count", "sum_value", "min_value", "max_value",
    ], field_rows)

if __name__ == "__main__":
    main()





